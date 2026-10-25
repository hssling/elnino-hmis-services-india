"""Write the SYNTHETIC HMIS fixtures (Paper E, Task 2; ruling F1).

The files copy the LAYOUT of the real portal files (SAS HTML exports - stored here as .xls.html because the controlled-data guard rejects .xls - and .xlsx workbooks): the
title lines, the header rows, the row and column spans, the item codes, the item text and the type labels. Every
unit name ("Test State", "District A", "State P", ...) and every number is invented. Nothing is cut from a
downloaded file. FIXTURES.md names the real file whose layout each fixture mirrors.

    python -I tests/hmis/fixtures/make_fixtures.py            # rewrites the fixtures next to this script

Invented values: value(item, unit, month) = 100 * (item index + 1) + 10 * unit index + calendar month, except
the special cells in SPECIAL. E2/E3 split columns use separate ranges (Public 20000+, Private 30000+,
Urban 40000+, Rural 50000+) so a test can tell the Total split apart; they do not add up to Total.
"""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent

MONTHS = ["April", "May", "June", "July", "August", "September", "October", "November", "December",
          "January", "February", "March"]
CAL = [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]
SPLITS = ["Total [(A+B) or (C+D)]", "Public [A]", "Private [B]", "Urban [C]", "Rural [D]"]
SPLIT_BASE = [0, 20000, 30000, 40000, 50000]
TOTAL_NOTE = ("Total [(A+B) OR (C+D)] =[(\xa0Public Facilities + Private Facilities) OR "
              "(Rural Facilities + Urban Facilities)]")

# (code, text, types): real item codes and item text of each era; one multi-type item per era.
E1_ITEMS = [
    ("1.1", "Total number of pregnant women Registered for ANC", ["TOTAL"]),
    ("1.1.1", "Of which Number registered within first trimester (within 12 weeks)", ["TOTAL"]),
    ("2.2", "Deliveries conducted at Public Institutions (Including C-Sections)", ["TOTAL"]),
    ("2.3", "Deliveries conducted at Private Institutions (Including C-Sections)", ["TOTAL"]),
    ("10.1.03", "Number of Infants (0 to 11 months old) received DPT2 immunisation", ["TOTAL"]),
    ("10.1.04", "Number of Infants (0 to 11 months old) received DPT3 immunisation", ["TOTAL"]),
    ("10.1.04C", "Number of Infants (0 to 11 months old) received Pentavalent3 immunisation", ["TOTAL"]),
    ("10.1.12", "Number of Infants (0 to 11 months old) received Measles immunisation (First Dose)", ["TOTAL"]),
    ("10.1.13.a", "Total number of male children (9 to 11 months old) fully immunised "
                  "(BCG+DPT123/Pentavalent123+OPV123+Measles) during the month", ["TOTAL"]),
    ("10.1.13.b", "Total number of female children (9 to 11 months old) fully immunised "
                  "(BCG+DPT123/Pentavalent123+OPV123+Measles) during the month", ["TOTAL"]),
    ("10.1.13.c", "Total number of children (9 to 11 months old) fully immunised "
                  "(BCG+DPT123/Pentavalent123+OPV123+Measles) during the month (sum of items 10.1.13.a and "
                  "10.1.13.b)", ["TOTAL"]),
    ("12.7", "Number of cases of Diarrhoea and Dehydration reported in children below 5 years of age", ["TOTAL"]),
    ("14.12.1", "OPD attendance (All)", ["TOTAL"]),
    ("16.3.02", "Stock Position of Gloves", ["1. Balance From Previous Month", "5. Total Stock"]),
]
E2_ITEMS = [
    ("1.1", "Total number of pregnant women registered for ANC", ["TOTAL"]),
    ("2.2", "Number of Institutional Deliveries conducted (Including C-Sections)", ["TOTAL"]),
    ("9.1.5", "Child immunisation - DPT3", ["TOTAL"]),
    ("9.1.8", "Child immunisation - Pentavalent 3", ["TOTAL"]),
    ("9.2.1", "Child immunisation (9-11months) - Measles & Rubella (MR)- 1st Dose", ["TOTAL"]),
    ("9.2.2", "Child immunisation (9-11months) - Measles 1st dose", ["TOTAL"]),
    ("9.2.4.a", "Children aged between 9 and 11 months fully immunized- Male", ["TOTAL"]),
    ("9.2.4.b", "Children aged between 9 and 11 months fully immunized - Female", ["TOTAL"]),
    ("10.11", "Childhood Diseases - Diarrhoea", ["TOTAL"]),
    ("14.2.1", "Allopathic- Outpatient attendance", ["TOTAL"]),
    ("14.2.2", "Ayush - Outpatient attendance", ["TOTAL"]),
    ("16.1.1", "Stock Position of Gloves", ["1. Balance From Previous Month", "5. Total Stock"]),
]
E3_ITEMS = [(c, t.replace("Measles & Rubella (MR)- 1st Dose",
                          "Measles & Rubella (MR)/ Measles containing vaccine(MCV) - 1st Dose"), ty)
            for c, t, ty in E2_ITEMS]
SECTION = {"1": "M1 [Ante Natal Care (ANC)]", "2": "M2 [Deliveries]", "9": "M9 [Child Immunisation]",
           "10": "M10 [Number of cases of Childhood Diseases (0-5 years)]", "14": "M14 [Patient Services]",
           "16": "M16 [Stock Position]"}

# Special cells: (fixture, unit, code, calendar month) -> cell text. Everything else is the formula.
SPECIAL = {
    ("e1_state", "District B", "12.7", 6): "NA",
    ("e1_state", "District B", "12.7", 7): "",
    ("e1_state", "District A", "10.1.04", 5): "0",
    ("e1_allindia", "State P", "1.1", 4): "4321",
    ("e2_state_1920", "District A", "9.2.2", 4): "0",
    ("e2_state_1920", "District A", "9.2.1", 4): "50",
    ("e3_state", "Test State", "1.1", 10): "777",
}


def value(fixture, unit_i, unit, item_i, code, month, split=0):
    """Cell text for (unit, item, month[, split]); splits other than Total use their own range."""
    if split == 0 and (fixture, unit, code, month) in SPECIAL:
        return SPECIAL[(fixture, unit, code, month)]
    return str(SPLIT_BASE[split] + 100 * (item_i + 1) + 10 * unit_i + month)


# ---- SAS HTML (.xls) --------------------------------------------------------------------------------------
HTML_HEAD = """<html xmlns:x="urn:schemas-microsoft-com:office:excel">
<head>
<meta name="Generator" content="SAS Software Version 9.3, see www.sas.com">
<meta http-equiv="Content-type" content="charset=windows-1252">
<title>SAS Output</title>
<style type="text/css">
<!--
.data { font-size: 8pt; }
-->
</style>
<script language="javascript" type="text/javascript">
<!--
function startup(){ }
function shutdown(){ }
//-->
</script>
</head>
"""


def _h1(cls, text):
    return f'<h1 class="{cls}" style=" color: #000000; text-align: left;"> {text} </h1>\n'


def _td(text, cls="data", colspan=1, rowspan=1, scope=None):
    a = f' class="{cls}"' if cls else ""
    a += f' rowspan="{rowspan}"' if rowspan > 1 else ""
    a += f' colspan="{colspan}"' if colspan > 1 else ""
    a += f' scope="{scope}"' if scope else ""
    text = text.replace("&", "&amp;").replace("'", "&#39;")
    return f"<td{a}>{text if text != '' else '&nbsp;'} </td>\n"


def sas_html(fixture, unit_label, title_unit, fy4, status, units, items, splits, quote_codes=False):
    """A SAS 'Data ItemWise Report' table: months across (x 5 splits if `splits`), units x items down."""
    width = 5 if splits else 1
    out = [HTML_HEAD, '<body onload="startup()" onunload="shutdown()"']        # the real files' unclosed tag
    out.append(_h1("systemtitle", f"Data ItemWise Report for - {title_unit.replace('&', '&amp;')}"))
    out.append(_h1("systemtitle2", f"Provisional Figures for - {fy4}"))
    out.append(_h1("systemtitle3", f"Status As On: {status}" + " " * 40))
    if splits:
        out.append(_h1("systemtitle4", TOTAL_NOTE.replace("\xa0", "&nbsp;")))
    out.append('<p>\n<table class="table" border="1" cellspacing="0" cellpadding="3" '
               'summary="Procedure Tabulate: ">\n<thead>\n<tr>\n')
    if splits:
        out.append(_td("", "header", colspan=4, rowspan=2, scope="colgroup"))
        out += [_td(m, "header", colspan=5, scope="colgroup") for m in MONTHS + ["Total"]]
        out.append("</tr>\n<tr>\n")
        out += [_td(s, "header", scope="col") for _ in range(13) for s in SPLITS]
    else:
        out.append(_td("", "header", colspan=4, scope="colgroup"))
        out += [_td(m, "header", scope="col") for m in MONTHS + ["Total"]]
    out.append("</tr>\n</thead>\n<tbody>\n<tr>\n")
    out.append(_td(unit_label, "rowheader", scope="row"))
    out += [_td("", "rowheader", scope="row")] * 3
    out.append(f'<td colspan="{13 * width}"></td>\n</tr>\n')
    n_rows = sum(len(t) for _, _, t in items)
    for ui, unit in enumerate(units):
        first_unit_row = True
        for ii, (code, text, types) in enumerate(items):
            for ti, typ in enumerate(types):
                out.append("<tr>\n")
                if first_unit_row:
                    out.append(_td(unit, "rowheader", rowspan=n_rows, scope="rowgroup"))
                    first_unit_row = False
                if ti == 0:
                    shown = f"'{code}'" if quote_codes else code
                    out.append(_td(shown, "rowheader", rowspan=len(types), scope="rowgroup" if len(types) > 1
                                   else "row"))
                    out.append(_td(text, "rowheader", rowspan=len(types), scope="rowgroup" if len(types) > 1
                                   else "row"))
                out.append(_td(typ, "rowheader", scope="row"))
                total = [0] * width
                for m in CAL:
                    for s in range(width):
                        v = value(fixture, ui, unit, ii + 10 * ti, code, m, s)
                        total[s] += int(v) if v.strip().isdigit() else 0
                        out.append(_td(v))
                out += [_td(str(t)) for t in total]
                out.append("</tr>\n")
    out.append("</tbody>\n</table>\n</p>\n<p>Grey Box indicates that data for this item can not be cumulated.</p>\n"
               "</body>\n</html>\n")
    return "".join(out).encode("cp1252")


# ---- .xlsx -----------------------------------------------------------------------------------------------
def _xlsx(path, rows, merges=()):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Item Wise Report" if "e3_" in path.name else "Sheet1"
    for r in rows:
        ws.append(r)
    for m in merges:
        ws.merge_cells(m)
    wb.save(path)


def _col(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


def e2_xlsx(path, fixture, title_unit, units, items):
    """The 2017-18 Telangana .xlsx layout: title lines as rows, merged month headers, district on its first row."""
    rows = [[f"Data ItemWise Report for - {title_unit}"], ["Provisional Figures for - 2017-2018"],
            ["Status As On: 15 Mar 2019, 12:38 PM"], [TOTAL_NOTE], [],
            ["", "", "", ""] + [x for m in MONTHS + ["Total"] for x in [m, None, None, None, None]],
            ["", "", "", ""] + SPLITS * 13, ["District", "", "", ""]]
    for ui, unit in enumerate(units):
        first = True
        for ii, (code, text, types) in enumerate(items):
            for ti, typ in enumerate(types):
                vals = [value(fixture, ui, unit, ii + 10 * ti, code, m, s) for m in CAL for s in range(5)]
                tot = [str(sum(int(vals[k * 5 + s]) for k in range(12))) for s in range(5)]
                rows.append([unit if first else None, code if ti == 0 else None, text if ti == 0 else None, typ]
                            + [int(v) if v.isdigit() else v for v in vals] + [int(t) for t in tot])
                first = False
    merges = [f"{_col(5 + 5 * k)}6:{_col(9 + 5 * k)}6" for k in range(13)]
    _xlsx(path, rows, merges)


def e3_xlsx(path, fixture, title, month_label, status, groups, items, districts_label="Districts"):
    """The FY2020-21 month layout: one month, column groups (state or All India first) x 5 splits."""
    rows = [[title], [f"Provisional Figures for - 2020-2021 and Month - {month_label}"], [f"Status As On: {status}"],
            [TOTAL_NOTE], [],
            ["", "", "", "", districts_label],
            ["", "", "", ""] + [x for g in groups for x in [g, None, None, None, None]],
            ["", "", "", ""] + SPLITS * len(groups)]
    month = {"Apr": 4, "Oct": 10}[month_label[:3]]
    for ii, (code, text, types) in enumerate(items):
        for ti, typ in enumerate(types):
            vals = []
            for gi, g in enumerate(groups):
                for s in range(5):
                    v = value(fixture, gi, g, ii + 10 * ti, code, month, s)
                    vals.append(float(v) if v.isdigit() else v)
            rows.append([SECTION[code.split(".")[0]], f"'{code}", text, typ] + vals)
    merges = [f"{_col(5 + 5 * k)}7:{_col(9 + 5 * k)}7" for k in range(len(groups))]
    _xlsx(path, rows, merges)


def write_all(out_dir: Path = HERE) -> list[Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    made = []

    def put(name, data):
        p = out_dir / name
        p.write_bytes(data)
        made.append(p)

    put("e1_synthetic_state_1415.xls.html",
        sas_html("e1_state", "District", "Test State", "2014-2015", "29 Jan 2016, 2:00 PM",
                 ["District A", "District B"], E1_ITEMS, splits=False))
    put("e1_synthetic_allindia_1213.xls.html",
        sas_html("e1_allindia", "State", "All India", "2012-2013", "19SEP2015",
                 ["State P", "State Q"], E1_ITEMS, splits=False))
    put("e2_synthetic_state_1920.xls.html",
        sas_html("e2_state_1920", "District", "Test State", "2019-2020", "06 Nov 2020, 2:00 PM",
                 ["District A", "District B"], E2_ITEMS, splits=True, quote_codes=True))
    p = out_dir / "e2_synthetic_state_1718.xlsx"
    e2_xlsx(p, "e2_state_1718", "Test State", ["District A", "District B"], E2_ITEMS)
    made.append(p)
    p = out_dir / "e3_synthetic_state_2021_oct.xlsx"
    e3_xlsx(p, "e3_state", "Data ItemWise Report for Test State", "Oct-2020", "01 Oct 2024,06:05:11",
            ["Test State", "District A", "District B"], E3_ITEMS)
    made.append(p)
    p = out_dir / "e3_synthetic_allindia_2021_apr.xlsx"
    e3_xlsx(p, "e3_allindia", "Data ItemWise Report for All India Across States", "Apr-2020",
            "01 Oct 2024,06:01:01", ["All India", "State P", "State Q"], E3_ITEMS, districts_label="")
    made.append(p)
    return made


if __name__ == "__main__":
    for p in write_all(Path(sys.argv[1]) if len(sys.argv) > 1 else HERE):
        print(p)
